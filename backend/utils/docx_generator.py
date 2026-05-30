import copy
from pathlib import Path
from docx import Document
from lxml import etree
from backend.config import OUTPUTS_DIR

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


class DocxGenerator:
    def generate(self, resume_data, template_path, output_filename, job_type):
        doc = Document(str(template_path))

        # Summary
        self._expand_bullets(doc, "[[Summary]]", resume_data.get("summary", []))

        # Technical Skills
        skills_placeholder = "[[Technical_Skills]]" if job_type == "gc" else "[[Technical Skills]]"
        self._replace_skills(doc, skills_placeholder, resume_data.get("technical_skills", []))

        # Experience sections (4 for contract/fulltime, 5 for gc)
        exp_count = 5 if job_type == "gc" else 4
        has_desc = job_type == "contract"
        has_env = job_type in ("contract", "gc")

        for i in range(1, exp_count + 1):
            # Bullets
            self._expand_bullets(doc, f"[[Exp{i}]]", resume_data.get(f"exp{i}", []))

            # Description (contract only)
            if has_desc:
                desc = resume_data.get(f"exp{i}_description", "")
                if desc:
                    self._replace_text(doc, f"[[Exp{i}_Description]]", desc)

            # Environment (contract and gc)
            if has_env:
                env = resume_data.get(f"exp{i}_env", "")
                if env:
                    self._replace_env(doc, f"[[Exp{i}_Env]]", env)

        output_path = OUTPUTS_DIR / output_filename
        doc.save(str(output_path))
        return output_path

    def _expand_bullets(self, doc, placeholder, bullets):
        if not bullets:
            return
        for para in doc.paragraphs:
            if placeholder in para.text:
                p_element = para._element
                self._set_paragraph_text(para, bullets[0])
                prev = p_element
                for text in bullets[1:]:
                    new_p = copy.deepcopy(p_element)
                    self._clear_runs(new_p)
                    self._add_text_run(new_p, text, p_element)
                    prev.addnext(new_p)
                    prev = new_p
                return

    def _replace_text(self, doc, placeholder, text):
        for para in doc.paragraphs:
            if placeholder in para.text:
                self._set_paragraph_text(para, text)
                return

    def _replace_skills(self, doc, placeholder, skill_lines):
        if not skill_lines:
            return
        for para in doc.paragraphs:
            if placeholder in para.text:
                p_element = para._element
                self._set_skills_line(para, skill_lines[0])
                prev = p_element
                for line in skill_lines[1:]:
                    new_p = copy.deepcopy(p_element)
                    self._clear_runs(new_p)
                    self._add_skills_line_to_element(new_p, line, p_element)
                    prev.addnext(new_p)
                    prev = new_p
                return

    def _replace_env(self, doc, placeholder, env_text):
        for para in doc.paragraphs:
            if placeholder in para.text:
                for run in para.runs:
                    run.text = ""
                if para.runs:
                    para.runs[0].text = "Environment: "
                    para.runs[0].bold = True
                    clean = env_text
                    if clean.lower().startswith("environment:"):
                        clean = clean[len("Environment:"):].strip()
                    run = para.add_run(clean)
                    run.bold = False
                    if para.runs[0].font.size:
                        run.font.size = para.runs[0].font.size
                else:
                    para.text = env_text
                return

    def _set_paragraph_text(self, para, text):
        if para.runs:
            para.runs[0].text = text
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.text = text

    def _set_skills_line(self, para, line):
        if ":" not in line:
            self._set_paragraph_text(para, line)
            return
        cat, items = line.split(":", 1)
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = f"{cat.strip()}:"
            para.runs[0].bold = True
            fs = para.runs[0].font.size
            r = para.add_run(f" {items.strip()}")
            r.bold = False
            if fs:
                r.font.size = fs
        else:
            para.text = line

    def _clear_runs(self, p_element):
        for r in p_element.findall(f"{W}r"):
            p_element.remove(r)

    def _add_text_run(self, p_element, text, template_p):
        new_r = etree.SubElement(p_element, f"{W}r")
        template_runs = template_p.findall(f"{W}r")
        if template_runs:
            rpr = template_runs[0].find(f"{W}rPr")
            if rpr is not None:
                new_r.insert(0, copy.deepcopy(rpr))
        new_t = etree.SubElement(new_r, f"{W}t")
        new_t.text = text
        if text and (text[0] == " " or text[-1] == " "):
            new_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

    def _add_skills_line_to_element(self, p_element, line, template_p):
        if ":" not in line:
            self._add_text_run(p_element, line, template_p)
            return
        cat, items = line.split(":", 1)
        template_runs = template_p.findall(f"{W}r")

        bold_r = etree.SubElement(p_element, f"{W}r")
        bold_rpr = etree.SubElement(bold_r, f"{W}rPr")
        etree.SubElement(bold_rpr, f"{W}b")
        etree.SubElement(bold_rpr, f"{W}bCs")
        if template_runs:
            rpr = template_runs[0].find(f"{W}rPr")
            if rpr is not None:
                for tag in (f"{W}sz", f"{W}szCs"):
                    el = rpr.find(tag)
                    if el is not None:
                        bold_rpr.append(copy.deepcopy(el))
        bold_t = etree.SubElement(bold_r, f"{W}t")
        bold_t.text = f"{cat.strip()}:"
        bold_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

        items_r = etree.SubElement(p_element, f"{W}r")
        items_rpr = etree.SubElement(items_r, f"{W}rPr")
        if template_runs:
            rpr = template_runs[0].find(f"{W}rPr")
            if rpr is not None:
                for tag in (f"{W}sz", f"{W}szCs"):
                    el = rpr.find(tag)
                    if el is not None:
                        items_rpr.append(copy.deepcopy(el))
        items_t = etree.SubElement(items_r, f"{W}t")
        items_t.text = f" {items.strip()}"
        items_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
